"""더미 문제 10개 × 서류 2개 + 참고 서류를 만든다.

설계 기준: 노션 4.2 '기술 스택 결정 & 샘플 데이터' 3장.
회사·사람 이름은 가상이고, 메일 주소는 예약 도메인(.example)을 쓴다.
실행: python3 scripts/samples/build.py  →  samples/ 아래에 만들어진다.
"""
from pathlib import Path
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from render import (Renderer, kdate, make_eml, photoify, quote_block_en,
                    quote_block_ko, save_scan_pdf, scanify)

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "samples"

# ---------------- 가상 회사·사람 ----------------
GT = "가온테라로지스(주)"
GT_EN = "GAONTERRA LOGIS CO., LTD."
KIM = "김도현 <dhkim@gaonterra.example>"
PARK = "박서진 <sjpark@gaonterra.example>"
LEE = "이하준 <hjlee@gaonterra.example>"
CORVANE_JUNG = "정민재 (Corvane Line Korea) <mj.jung@corvane.example>"
TALVIK_HAN = "Grace Han (Talvik Container Line Korea) <grace.han@talvik.example>"
MIRSOL_CHOI = "최유나 (미르솔텍) <yuna.choi@mirsoltech.example>"
HANUI_HAN = "한지수 (하늬소재) <jshan@hanui.example>"
MIRSOL_BUY = "미르솔텍 구매팀 <purchase@mirsoltech.example>"

XLSX = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# =========================================================
# HTML 틀
# =========================================================
def gt_quote(no, date, to, owner, route, term, cntr, validity, items, notes, total=None,
             big=None, notes_bottom=False, extra_css=""):
    """가온테라로지스 견적서. items: [(항목, 단위, 금액)]"""
    rows = "".join(
        f"<tr><td>{a}</td><td class='c'>{b}</td><td class='r'>{c}</td></tr>" for a, b, c in items)
    total_row = f"<tr><th colspan='2' class='r'>합계</th><th class='r'>{total}</th></tr>" if total else ""
    big_html = ""
    if big:
        big_html = f"""
        <div style='margin:26px 0 8px;border:2px solid #1d3557;padding:18px 24px;display:flex;
             justify-content:space-between;align-items:center'>
          <div style='font-size:22px;font-weight:700'>{big[0]}</div>
          <div style='font-size:46px;font-weight:800;letter-spacing:1px'>{big[1]}</div>
        </div>"""
    note_html = f"<div class='small' style='line-height:1.6'>{notes}</div>"
    note_pos = (f"<div style='position:absolute;left:60px;right:60px;bottom:60px'>{note_html}</div>"
                if notes_bottom else f"<div style='margin-top:18px'>{note_html}</div>")
    return f"""
<div class='page' style='{extra_css}'>
  <div style='display:flex;justify-content:space-between;align-items:flex-end;border-bottom:3px solid #1d3557;padding-bottom:12px'>
    <div>
      <div style='font-size:24px;font-weight:800;color:#1d3557'>{GT}</div>
      <div class='small muted'>{GT_EN} · 부산광역시 · TEL 051-000-2814</div>
    </div>
    <div style='font-size:32px;font-weight:800;letter-spacing:14px'>견 적 서</div>
  </div>
  <table style='margin-top:22px'>
    <tr><th style='width:16%'>견적번호</th><td style='width:34%'>{no}</td><th style='width:16%'>견적일</th><td>{date}</td></tr>
    <tr><th>수신</th><td>{to}</td><th>담당</th><td>{owner}</td></tr>
    <tr><th>구간</th><td colspan='3'>{route}</td></tr>
    <tr><th>조건</th><td>{term}</td><th>컨테이너</th><td>{cntr}</td></tr>
    <tr><th>유효기간</th><td colspan='3'>{validity}</td></tr>
  </table>
  {big_html}
  <table style='margin-top:18px'>
    <tr><th style='width:52%'>항목</th><th style='width:18%'>단위</th><th>금액</th></tr>
    {rows}{total_row}
  </table>
  <div style='margin-top:26px' class='small muted'>위와 같이 견적합니다. 감사합니다.</div>
  {note_pos}
</div>"""


def carrier_header(name, sub, color):
    return f"""
  <div style='background:{color};color:#fff;padding:16px 22px;display:flex;justify-content:space-between;align-items:center'>
    <div style='font-size:24px;font-weight:800;letter-spacing:2px'>{name}</div>
    <div class='small' style='text-align:right;line-height:1.5'>{sub}</div>
  </div>"""


# =========================================================
# 엑셀 도구
# =========================================================
THIN = Side(style="thin", color="555555")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E7E7E7")


def xl_cell(ws, ref, value, bold=False, fill=None, border=True, align=None, size=10, num=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name="맑은 고딕", bold=bold, size=size)
    if fill:
        c.fill = fill
    if border:
        c.border = BOX
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if num:
        c.number_format = num
    return c


def wb_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =========================================================
# 문제별 서류
# =========================================================
def q01(r: Renderer, d: Path):
    """① 인상 이유 설명: 지난 견적(사진) ↔ 새 운임(메일 2통). 부산→로테르담 40ft 1대."""
    html = gt_quote(
        "GT-2609-031", "2026. 09. 08", "윤슬정밀 귀중 (구매팀)", "김도현",
        "부산 (BUSAN, KOREA) → 로테르담 (ROTTERDAM, NETHERLANDS)", "CY – CY", "40' × 1",
        "2026. 09. 01 ~ 2026. 09. 30",
        [("Ocean Freight", "40'", "USD 1,800")],
        "비고: THC 포함, 서류비 등 기타 비용 별도 / 상기 운임은 선사 사정에 따라 변동될 수 있습니다.",
        total="USD 1,800", big=("Ocean Freight", "USD 1,800"), notes_bottom=True)
    img = r.png(html)
    img.save(d / "A_지난견적_9월_인쇄원본.png")
    r.pdf(html, d / "A_지난견적_9월_인쇄원본.pdf")
    photoify(img, seed=11).save(d / "A_지난견적_9월_사진(합성).jpg", quality=82)

    req_body = ("정민재 차장님, 안녕하세요.\n가온테라로지스 김도현입니다.\n\n"
                "10월 선적분 부산 → 로테르담 40' 1대 운임 부탁드립니다.\n\n감사합니다.\n김도현 드림")
    body = ("김도현 대표님, 안녕하세요.\nCorvane Line 정민재입니다.\n\n"
            "요청하신 10월 운임 아래와 같이 회신드립니다.\n\n"
            "POL: BUSAN, KOREA (KRPUS)\nPOD: ROTTERDAM, NETHERLANDS (NLRTM)\n"
            "CNTR: 40' x 1\nO/F USD 1,700/40', CY/CY\nLocal charges as per tariff\n"
            "Validity: Oct 1 - Oct 31, 2026\n\n"
            "감사합니다.\n정민재 드림\nCorvane Line Korea | Busan Sales")
    body += quote_block_ko(KIM, CORVANE_JUNG, kdate(2026, 10, 2, 9, 40),
                           "[가온테라] 10월 부산-로테르담 40' 운임 요청", req_body)
    make_eml(d / "B_선사회신_10월.eml", CORVANE_JUNG, KIM,
             "RE: [가온테라] 10월 부산-로테르담 40' 운임 요청", kdate(2026, 10, 5, 14, 20), body)


def q02(r: Renderer, d: Path):
    """② 지난 견적(엑셀, 8월) ↔ 새 운임(글자 PDF)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "견적"
    for col, w in zip("ABCDEF", [18, 26, 10, 12, 8, 14]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:F1")
    xl_cell(ws, "A1", f"{GT} 견적서", bold=True, size=14, border=False, align="center")
    meta = [("견적일", "2026-08-11"), ("수신", "윤슬정밀 (구매팀)"), ("담당", "김도현"),
            ("구간", "부산(BUSAN) → 함부르크(HAMBURG)"), ("조건", "CY - CY"),
            ("컨테이너", "40'HC × 1"), ("유효기간", "2026-08-01 ~ 2026-08-31")]
    for i, (k, v) in enumerate(meta, start=3):
        xl_cell(ws, f"A{i}", k, bold=True, fill=HEAD_FILL)
        ws.merge_cells(f"B{i}:F{i}")
        xl_cell(ws, f"B{i}", v)
    hr = 11
    for col, h in zip("ABCDEF", ["항목", "내용", "통화", "단가", "수량", "금액"]):
        xl_cell(ws, f"{col}{hr}", h, bold=True, fill=HEAD_FILL, align="center")
    xl_cell(ws, "A12", "All-in")
    xl_cell(ws, "B12", "해상운임 (40'HC)")
    xl_cell(ws, "C12", "USD", align="center")
    xl_cell(ws, "D12", 2000, num="#,##0")
    xl_cell(ws, "E12", 1, align="center")
    xl_cell(ws, "F12", 2000, num="#,##0")
    xl_cell(ws, "A13", "합계", bold=True)
    ws.merge_cells("B13:E13")
    xl_cell(ws, "B13", "")
    xl_cell(ws, "F13", 2000, bold=True, num="#,##0")
    # 20줄 아래 각주
    xl_cell(ws, "A33", "※ All-in 운임은 기본운임(O/F)·유류할증료(BAF) 포함, 터미널 조작료(THC)·서류비 별도입니다.",
            border=False, size=8)
    wb.save(d / "A_지난견적_8월.xlsx")

    html = f"""
<div class='page'>
  {carrier_header("SELVARO SHIPPING", "Selvaro Shipping Korea<br>Busan Branch", "#1f6f5c")}
  <div style='font-size:22px;font-weight:800;margin:26px 0 6px'>QUOTATION</div>
  <table>
    <tr><th style='width:22%'>To</th><td>GAONTERRA LOGIS CO., LTD. (Mr. Kim)</td><th style='width:18%'>Date</th><td>2026-10-06</td></tr>
    <tr><th>POL</th><td>BUSAN, KOREA (KRPUS)</td><th>POD</th><td>HAMBURG, GERMANY (DEHAM)</td></tr>
    <tr><th>Equipment</th><td>40'HC x 1</td><th>Term</th><td>CY/CY</td></tr>
    <tr><th>Validity</th><td colspan='3'>2026-10-01 ~ 2026-10-31</td></tr>
  </table>
  <table style='margin-top:20px'>
    <tr><th style='width:46%'>Charge</th><th style='width:20%'>Unit</th><th>Rate</th></tr>
    <tr><td>Basic Ocean Freight (O/F)</td><td class='c'>40'HC</td><td class='r'>USD 1,650</td></tr>
    <tr><td>Bunker Surcharge (BAF)</td><td class='c'>40'HC</td><td class='r'>as per monthly notice</td></tr>
    <tr><td>Origin THC</td><td class='c'>40'HC</td><td class='r'>as per local tariff</td></tr>
  </table>
  <div class='small muted' style='margin-top:16px;line-height:1.6'>
    Subject to space and equipment availability.<br>
    Rates are subject to change with prior notice.
  </div>
</div>"""
    r.pdf(html, d / "B_선사견적_10월.pdf")


def q03(r: Renderer, d: Path):
    """③ 지난 견적(스캔 PDF, 원화) ↔ 새 운임(메일)."""
    html = gt_quote(
        "GT-2609-044", "2026. 09. 10", "다래하임 귀중 (물류팀)", "김도현",
        "부산 (BUSAN, KOREA) → 롱비치 (LONG BEACH, USA)", "CY – CY", "40'HC × 1",
        "2026. 09. 01 ~ 2026. 09. 30",
        [("해상운임", "40'HC", "2,484,000원")],
        "※ 적용 환율 1 USD = 1,380원 (9월 기준) / THC 별도",
        total="2,484,000원")
    scan = scanify(r.png(html), seed=31, angle=0.9)
    save_scan_pdf(scan, d / "A_지난견적_9월_스캔.pdf")

    body = ("Dear Mr. Kim,\n\nThank you for your inquiry.\nPlease find our October offer below.\n\n"
            "POL: Busan (KRPUS)\nPOD: Long Beach (USLGB)\nTerm: CY/CY\n"
            "40'HC O/F USD 1,760, THC excl.\nValidity: until Oct 31, 2026\n\n"
            "Best regards,\nGrace Han\nTalvik Container Line Korea")
    make_eml(d / "B_선사회신_10월.eml", TALVIK_HAN, KIM,
             "RE: Oct rate request PUS-LGB 40HC", kdate(2026, 10, 7, 11, 5), body)


def q04(r: Renderer, d: Path):
    """④ 대조 문제: 지난 견적(글자 PDF) ↔ 새 운임(엑셀). 함정 없음."""
    html = gt_quote(
        "GT-2609-052", "2026. 09. 15", "다래하임 귀중 (물류팀)", "김도현",
        "부산 (BUSAN, KOREA) → 싱가포르 (SINGAPORE)", "CY – CY", "40' × 1",
        "2026. 09. 01 ~ 2026. 09. 30",
        [("O/F (THC 포함)", "40'", "USD 520")],
        "상기 운임은 선사 사정에 따라 변동될 수 있습니다.", total="USD 520")
    r.pdf(html, d / "A_지난견적_9월.pdf")

    wb = Workbook()
    ws = wb.active
    ws.title = "Rate"
    ws.merge_cells("A1:G1")
    xl_cell(ws, "A1", "CORVANE LINE KOREA - INTRA ASIA RATE (OCT 2026)", bold=True, size=12, border=False)
    xl_cell(ws, "A2", "To: GAONTERRA LOGIS / Date: 2026-10-06", border=False, size=9)
    heads = ["POL", "POD", "Equipment", "O/F (USD)", "THC", "Term", "Validity"]
    for col, h in zip("ABCDEFG", heads):
        xl_cell(ws, f"{col}4", h, bold=True, fill=HEAD_FILL, align="center")
    vals = ["BUSAN (KRPUS)", "SINGAPORE (SGSIN)", "40'", 580, "THC 포함", "CY/CY", "2026-10-01 ~ 2026-10-31"]
    for col, v in zip("ABCDEFG", vals):
        xl_cell(ws, f"{col}5", v, num="#,##0" if isinstance(v, int) else None)
    for col, w in zip("ABCDEFG", [16, 20, 11, 11, 11, 9, 24]):
        ws.column_dimensions[col].width = w
    wb.save(d / "B_선사운임_10월.xlsx")


def q05(r: Renderer, d: Path):
    """⑤ 견적 회신: 화주 문의(메일) ↔ 운임표(글자 PDF). 3.3의 3-2."""
    body = ("박서진 과장님, 안녕하세요.\n미르솔텍 최유나입니다.\n\n"
            "11월 둘째 주 선적 예정인 건 견적 부탁드립니다.\n\n"
            "- 구간: 부산 CY → 로테르담 (도착항 비용까지 포함해서 받고 싶습니다)\n"
            "- 컨테이너: 40ft 1대\n- 품목: 산업용 필터 부품 (일반화물)\n"
            "- 선적 예정: 11월 둘째 주\n\n"
            "회신 주시면 내부 보고 후 바로 부킹 요청드리겠습니다.\n감사합니다.\n\n최유나 드림")
    make_eml(d / "A_화주문의.eml", MIRSOL_CHOI, PARK,
             "견적 요청 드립니다 (부산→로테르담, 40ft)", kdate(2026, 10, 27, 10, 12), body)

    html = f"""
<div class='page'>
  {carrier_header("CORVANE LINE", "Corvane Line Korea<br>Asia - North Europe Service", "#163e73")}
  <div style='font-size:20px;font-weight:800;margin:26px 0 4px'>OCEAN FREIGHT RATE SHEET</div>
  <div class='muted' style='margin-bottom:14px'>Asia → North Europe · Issued 2026-10-01 · To: GAONTERRA LOGIS</div>
  <table>
    <tr><th>POL</th><th>POD</th><th>Charge</th><th style='width:16%'>40'</th><th style='width:14%'>Term</th></tr>
    <tr><td rowspan='3'>BUSAN<br>(KRPUS)</td><td rowspan='3'>ROTTERDAM<br>(NLRTM)</td>
        <td>Ocean Freight</td><td class='r'>USD 1,700</td><td rowspan='3' class='c'>CY/CY</td></tr>
    <tr><td>Origin THC</td><td class='r'>USD 200</td></tr>
    <tr><td>DTHC (Destination THC)</td><td class='r'>TBA</td></tr>
  </table>
  <div style='font-size:8px;color:#444;margin-top:4px'>Valid till 2026-10-31</div>
  <div class='small muted' style='margin-top:30px;line-height:1.6'>
    Subject to space and equipment availability.<br>
    Cargo must be general cargo. DG / OOG on request.
  </div>
</div>"""
    r.pdf(html, d / "B_선사운임표.pdf")


def q06(r: Renderer, d: Path):
    """⑥ 견적 회신: 화주 문의(메일 스레드 4통) ↔ 선사 운임(엑셀)."""
    m1 = ("박서진 과장님,\n9월 말 호치민 건 40ft 1대 CY 기준 운임 부탁드립니다.\n감사합니다.\n한지수 드림")
    m2 = ("한지수 과장님,\n호치민(Cat Lai) CY 기준 40ft 1대 USD 520입니다.\n"
          "(해상운임·THC 포함, 9월 선적분)\n감사합니다.\n박서진 드림")
    m3 = ("과장님, 지난번 호치민 건 잘 진행됐습니다. 감사합니다.\n"
          "이번엔 하이퐁으로 40ft 1대 나갈 예정인데 운임 다시 부탁드려요.\n한지수 드림")
    m4 = ("박서진 과장님, 안녕하세요.\n하늬소재 한지수입니다.\n\n"
          "지난주 말씀드린 부산 → 하이퐁 40ft 1대 건 관련해서 정리해서 다시 보내드립니다.\n\n"
          "출항은 11월 3일 희망하고, 화물은 10월 30일까지 준비됩니다.\n"
          "품목은 지난번과 같은 섬유 소재(롤 원단)이고, 포장은 팔레트 18개입니다.\n"
          "바이어 요청으로 이번엔 하노이 공장 문 앞까지 보내 주셔야 할 것 같습니다.\n"
          "인보이스와 패킹리스트는 다음 주 초에 보내드리겠습니다.\n"
          "혹시 팔레트 높이 제한이 있으면 미리 알려 주세요.\n\n"
          "바쁘시겠지만 확인 부탁드립니다.\n감사합니다.\n\n한지수 드림")
    body = m4
    body += quote_block_ko(HANUI_HAN, PARK, kdate(2026, 10, 20, 16, 30), "RE: RE: 호치민 운임 문의", m3)
    body += quote_block_ko(PARK, HANUI_HAN, kdate(2026, 9, 16, 9, 50), "RE: 호치민 운임 문의", m2)
    body += quote_block_ko(HANUI_HAN, PARK, kdate(2026, 9, 15, 15, 2), "호치민 운임 문의", m1)
    make_eml(d / "A_화주문의_스레드.eml", HANUI_HAN, PARK,
             "RE: RE: RE: 호치민 운임 문의", kdate(2026, 10, 27, 13, 45), body)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rate"
    ws.merge_cells("A1:I1")
    xl_cell(ws, "A1", "SELVARO SHIPPING - INTRA ASIA RATE", bold=True, size=12, border=False)
    xl_cell(ws, "A2", "To: GAONTERRA LOGIS / Issued: 2026-10-14 / Currency: USD", border=False, size=9)
    heads = ["POL", "POD", "20'", "40'", "THC 20'", "THC 40'", "Term", "Validity", "Remark"]
    for col, h in zip("ABCDEFGHI", heads):
        xl_cell(ws, f"{col}4", h, bold=True, fill=HEAD_FILL, align="center")
    rows = [
        ["BUSAN", "HAIPHONG (VNHPH)", 280, 480, 150, 200, "CY/CY", "2026-10-15 ~ 2026-11-15", ""],
        ["BUSAN", "HO CHI MINH (VNSGN)", 300, 520, 150, 200, "CY/CY", "2026-10-15 ~ 2026-11-15", "Cat Lai"],
        ["BUSAN", "BANGKOK (THBKK)", 320, 560, 150, 200, "CY/CY", "2026-10-15 ~ 2026-11-15", "Laem Chabang +50"],
    ]
    for i, row in enumerate(rows, start=5):
        for col, v in zip("ABCDEFGHI", row):
            xl_cell(ws, f"{col}{i}", v, num="#,##0" if isinstance(v, int) else None)
    for col, w in zip("ABCDEFGHI", [10, 22, 8, 8, 9, 9, 8, 24, 18]):
        ws.column_dimensions[col].width = w
    wb.save(d / "B_선사운임.xlsx")


def q07(r: Renderer, d: Path):
    """⑦ 견적 회신: 선적요청서(스캔) ↔ 선사 회신(메일 + 첨부 엑셀)."""
    html = """
<div class='page' style='font-family:"Noto Serif CJK KR",serif'>
  <div class='small' style='display:flex;justify-content:space-between;border-bottom:1px solid #000;padding-bottom:4px'>
    <span>2026-10-26 15:42</span><span>FAX 051-000-7731</span><span>별마루테크</span><span>P.1/1</span>
  </div>
  <div style='text-align:center;font-size:28px;font-weight:700;margin:30px 0 4px;letter-spacing:8px'>선 적 요 청 서</div>
  <div style='text-align:center;margin-bottom:22px'>SHIPPING REQUEST</div>
  <table>
    <tr><th style='width:22%'>수신</th><td>가온테라로지스 박서진 과장님</td></tr>
    <tr><th>발신</th><td>별마루테크 수출팀 문태오</td></tr>
    <tr><th>요청일</th><td>2026. 10. 26</td></tr>
  </table>
  <table style='margin-top:16px'>
    <tr><th style='width:22%'>SHIPPER</th><td>별마루테크</td></tr>
    <tr><th>CONSIGNEE</th><td>TO ORDER</td></tr>
    <tr><th>NOTIFY</th><td>SAME AS CONSIGNEE</td></tr>
    <tr><th>POL / POD</th><td>BUSAN, KOREA / LOS ANGELES, USA</td></tr>
    <tr><th>CONTAINER</th><td>20' × 2 (CY)</td></tr>
    <tr><th>ETD</th><td>2026. 11. 12</td></tr>
    <tr><th>COMMODITY</th><td>LED 조명 부품 (일반화물)</td></tr>
    <tr><th>Q'TY / G.W.</th><td>36 PLT / 18,400 KG</td></tr>
  </table>
  <div style='margin-top:20px;line-height:1.8'>
    위와 같이 선적을 요청드리오니 운임 견적과 스케줄 회신 부탁드립니다.<br>
    서류(인보이스·패킹리스트)는 확정 후 별도 송부하겠습니다.
  </div>
  <div style='margin-top:40px;text-align:right'>별마루테크 수출팀 문태오 (인)</div>
</div>"""
    scan = scanify(r.png(html), seed=71, angle=-1.1)
    save_scan_pdf(scan, d / "A_선적요청서_스캔.pdf")

    wb = Workbook()
    ws = wb.active
    ws.title = "USWC NOV"
    ws.merge_cells("A1:G1")
    xl_cell(ws, "A1", "TALVIK CONTAINER LINE - US WEST COAST RATE (NOV 2026)", bold=True, size=12, border=False)
    xl_cell(ws, "A2", "POL: BUSAN / Currency: USD / Term: CY/CY / THC excl. / Valid: 2026-11-01 ~ 2026-11-30",
            border=False, size=9)
    heads = ["POD", "UN/LOCODE", "20'", "40'", "40'HC", "T/T (days)", "Remark"]
    for col, h in zip("ABCDEFG", heads):
        xl_cell(ws, f"{col}4", h, bold=True, fill=HEAD_FILL, align="center")
    rows = [
        ["LOS ANGELES", "USLAX", "N/A", 1600, 1700, 13, ""],
        ["LONG BEACH", "USLGB", "N/A", 1600, 1700, 13, ""],
        ["OAKLAND", "USOAK", "N/A", 1700, 1800, 15, ""],
        ["SEATTLE", "USSEA", "N/A", 1750, 1850, 12, ""],
        ["VANCOUVER", "CAVAN", "N/A", 1800, 1900, 12, "Canada"],
    ]
    for i, row in enumerate(rows, start=5):
        for col, v in zip("ABCDEFG", row):
            xl_cell(ws, f"{col}{i}", v, num="#,##0" if isinstance(v, int) and col in "DE" else None,
                    align="center" if v == "N/A" else None)
    for col, w in zip("ABCDEFG", [14, 12, 8, 8, 8, 10, 10]):
        ws.column_dimensions[col].width = w
    att = wb_bytes(wb)
    att_name = "Talvik_USWC_Nov2026.xlsx"
    (d / f"B_첨부_{att_name}").write_bytes(att)

    body = ("박서진 과장님, 안녕하세요.\nTalvik Container Line 한그레이스입니다.\n\n"
            "요청하신 미주 서안 11월 운임 첨부드립니다.\n확인 부탁드립니다.\n\n"
            "감사합니다.\nGrace Han\nTalvik Container Line Korea")
    make_eml(d / "B_선사회신.eml", TALVIK_HAN, PARK, "RE: 미주 서안 11월 운임 요청",
             kdate(2026, 10, 27, 9, 30), body, attachments=[(att_name, att, *XLSX)])


def q08(r: Renderer, d: Path):
    """⑧ 재협상: 우리 견적(엑셀) ↔ 다른 포워더 견적(사진). 3.3의 3-3."""
    wb = Workbook()
    ws = wb.active
    ws.title = "견적"
    for col, w in zip("ABCDE", [22, 30, 10, 10, 14]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:E1")
    xl_cell(ws, "A1", f"{GT} 견적서", bold=True, size=14, border=False, align="center")
    meta = [("견적일", "2026-10-22"), ("수신", "가람빛화학 (물류 담당)"), ("담당", "이하준"),
            ("구간", "화주 공장(경남 김해) 상차 → 부산항 → 로테르담항"), ("조건", "Door - CY (공장 상차부터)"),
            ("컨테이너", "40' × 1"), ("유효기간", "2026-11-01 ~ 2026-11-30")]
    for i, (k, v) in enumerate(meta, start=3):
        xl_cell(ws, f"A{i}", k, bold=True, fill=HEAD_FILL)
        ws.merge_cells(f"B{i}:E{i}")
        xl_cell(ws, f"B{i}", v)
    hr = 11
    for col, h in zip("ABCDE", ["항목", "내용", "통화", "단위", "금액"]):
        xl_cell(ws, f"{col}{hr}", h, bold=True, fill=HEAD_FILL, align="center")
    items = [("해상운임 (O/F)", "부산항 → 로테르담항", "USD", "40'", 1700),
             ("THC", "부산항 터미널 조작료", "USD", "40'", 200),
             ("내륙운송", "김해 공장 → 부산신항", "USD", "40'", 350),
             ("수수료", "Handling Charge", "USD", "건", 50)]
    for i, it in enumerate(items, start=hr + 1):
        for col, v in zip("ABCDE", it):
            xl_cell(ws, f"{col}{i}", v, num="#,##0" if isinstance(v, int) else None)
    tr = hr + 1 + len(items)
    xl_cell(ws, f"A{tr}", "합계", bold=True)
    ws.merge_cells(f"B{tr}:D{tr}")
    xl_cell(ws, f"B{tr}", "")
    xl_cell(ws, f"E{tr}", 2300, bold=True, num="#,##0")
    wb.save(d / "A_우리견적.xlsx")

    # 상대 견적: 아래 작은 글씨 마지막 줄이 종이 끝에서 잘린다
    html = """
<div class='page' style='padding:70px 64px;overflow:hidden'>
  <div style='display:flex;align-items:center;gap:14px;border-bottom:2px solid #0b5fa5;padding-bottom:12px'>
    <div style='width:42px;height:42px;border-radius:50%;background:#0b5fa5'></div>
    <div><div style='font-size:22px;font-weight:800;color:#0b5fa5'>블루펜로지스</div>
    <div class='small muted'>BLUEFEN LOGIS · International Freight Forwarder</div></div>
  </div>
  <div style='font-size:18px;margin:34px 0 10px;color:#333'>Freight Quotation</div>
  <div class='small muted'>To: 가람빛화학 귀중 · Date: 2026.10.24</div>
  <div style='margin:46px 0;text-align:center'>
    <div style='font-size:20px;color:#333'>OCEAN FREIGHT</div>
    <div style='font-size:72px;font-weight:800;color:#0b5fa5;line-height:1.1'>USD 1,500</div>
    <div style='font-size:16px;color:#555'>per 40' container</div>
  </div>
  <div style='position:absolute;left:64px;bottom:90px;font-size:10.5px;color:#333;line-height:1.7;white-space:nowrap'>
    Route: Busan Port (CY) → Rotterdam Port (CY) / Equipment: 40' x 1<br>
    Valid until 2026.11.15 / Payment: Prepaid<br>
    * 상기 금액은 해상운임 기준이며, 부산 현지 부대비용(THC 등) 및 내륙운송비 포함 여부는 화물 조건 확인 후 별도로 안내해 드립니다
  </div>
</div>"""
    img = r.png(html)
    # 인쇄본도 마지막 줄이 오른쪽 끝에서 잘리도록, 오른쪽 일부를 잘라 A4에 얹는다
    w, h = img.size
    cut = img.crop((0, 0, int(w * 0.80), h))
    page = img.copy()
    page.paste((255, 255, 255), (0, 0, w, h))
    page.paste(cut, (0, 0))
    page.save(d / "B_다른포워더견적_인쇄원본.png")
    page.convert("RGB").save(d / "B_다른포워더견적_인쇄원본.pdf", "PDF", resolution=192.0)
    photoify(page, seed=81, crop_right=0.06).save(d / "B_다른포워더견적_사진(합성).jpg", quality=80)


def q09(r: Renderer, d: Path):
    """⑨ 재협상: 우리 견적(메일 스레드) ↔ 다른 포워더 견적(사진). 상대 유효기간이 지남."""
    old = ("최유나 대리님, 안녕하세요.\n가온테라로지스 이하준입니다.\n\n"
           "10월분 부산 → 앤트워프 40ft 1대 견적 드립니다.\n"
           "- USD 1,850 (해상운임·THC 포함, CY-CY)\n- 유효기간: 2026-10-01 ~ 10-31\n\n감사합니다.\n이하준 드림")
    new = ("최유나 대리님, 안녕하세요.\n가온테라로지스 이하준입니다.\n\n"
           "요청하신 11월분 부산 → 앤트워프 40ft 1대 견적 드립니다.\n"
           "- USD 1,980 (해상운임·THC 포함, CY-CY)\n- 유효기간: 2026-11-01 ~ 11-30\n\n"
           "선사 11월 운임 인상분이 반영된 금액입니다.\n감사합니다.\n이하준 드림")
    body = new + quote_block_ko(LEE, MIRSOL_CHOI, kdate(2026, 9, 18, 11, 20),
                                "[가온테라] 10월 부산-앤트워프 견적", old)
    make_eml(d / "A_우리견적_스레드.eml", LEE, MIRSOL_CHOI,
             "RE: [가온테라] 11월 부산-앤트워프 견적", kdate(2026, 10, 21, 15, 10), body)

    html = """
<div class='page' style='padding:64px'>
  <div style='text-align:right;font-size:12px;color:#666'>No. NW-2609-118</div>
  <div style='font-size:26px;font-weight:800;color:#2c3e50;margin-top:6px'>노르웰국제물류</div>
  <div class='small muted'>NORWELL GLOBAL LOGISTICS</div>
  <div style='margin-top:40px;font-size:20px;font-weight:700'>운임 견적 (Quotation)</div>
  <table style='margin-top:16px'>
    <tr><th style='width:26%'>수신</th><td>미르솔텍 귀중</td></tr>
    <tr><th>구간</th><td>Busan CY → Antwerp CY</td></tr>
    <tr><th>컨테이너</th><td>40' × 1</td></tr>
    <tr><th>포함 항목</th><td>O/F incl. THC</td></tr>
  </table>
  <div style='margin:50px 0;text-align:center'>
    <div style='font-size:64px;font-weight:800;color:#2c3e50'>USD 1,800</div>
    <div style='color:#555'>per container</div>
  </div>
  <div style='position:absolute;left:64px;bottom:70px;font-size:10px;color:#444'>Valid until 2026.10.15 · 상기 운임은 선사 사정에 따라 변동될 수 있습니다.</div>
</div>"""
    img = r.png(html)
    img.save(d / "B_다른포워더견적_인쇄원본.png")
    r.pdf(html, d / "B_다른포워더견적_인쇄원본.pdf")
    photoify(img, seed=91).save(d / "B_다른포워더견적_사진(합성).jpg", quality=80)


def q10(r: Renderer, d: Path):
    """⑩ 재협상: 우리 견적(글자 PDF) ↔ 다른 포워더 견적(스캔, 원화·건당)."""
    html = gt_quote(
        "GT-2610-077", "2026. 10. 22", "별마루테크 귀중 (수출팀)", "이하준",
        "부산 (BUSAN, KOREA) → 제벨알리 (JEBEL ALI, UAE)", "CY – CY", "40' × 2",
        "2026. 11. 01 ~ 2026. 11. 30",
        [("해상운임 (O/F)", "USD 1,700 / 컨테이너 × 2", "USD 3,400"),
         ("THC", "USD 200 / 컨테이너 × 2", "USD 400"),
         ("서류비 (DOC)", "USD 50 / BL", "USD 50")],
        "참고 환율 1 USD = 1,400원", total="USD 3,850")
    r.pdf(html, d / "A_우리견적.pdf")

    html2 = """
<div class='page' style='font-family:"Noto Serif CJK KR",serif'>
  <div class='small' style='display:flex;justify-content:space-between;border-bottom:1px solid #000;padding-bottom:4px'>
    <span>2026-10-23 10:18</span><span>FAX 051-000-4402</span><span>소담해로해운항공</span><span>P.1/1</span>
  </div>
  <div style='font-size:24px;font-weight:700;margin:28px 0 2px'>소담해로해운항공(주)</div>
  <div class='small'>SODAM HAERO SHIPPING &amp; AIR</div>
  <div style='text-align:center;font-size:26px;font-weight:700;margin:26px 0;letter-spacing:10px'>견 적 서</div>
  <table>
    <tr><th style='width:20%'>수신</th><td>별마루테크 수출팀 귀하</td><th style='width:16%'>일자</th><td>2026. 10. 23</td></tr>
    <tr><th>구간</th><td colspan='3'>부산항 → 제벨알리항 (UAE)</td></tr>
    <tr><th>물량</th><td colspan='3'>40ft 2대 기준</td></tr>
    <tr><th>유효기간</th><td colspan='3'>2026. 11. 30 까지</td></tr>
  </table>
  <table style='margin-top:18px'>
    <tr><th style='width:46%'>항목</th><th style='width:20%'>기준</th><th>금액 (원)</th></tr>
    <tr><td>해상운임</td><td class='c'>건당</td><td class='r'>4,620,000</td></tr>
    <tr><td>터미널 화물처리비</td><td class='c'>건당</td><td class='r'>532,000</td></tr>
    <tr><td>서류발급비</td><td class='c'>건당</td><td class='r'>70,000</td></tr>
    <tr><th colspan='2' class='r'>합계</th><th class='r'>5,222,000</th></tr>
  </table>
  <div style='margin-top:18px;font-size:12px;line-height:1.8'>
    ※ 부가세 별도 / 상기 금액은 40ft 2대 전체 금액입니다.
  </div>
</div>"""
    scan = scanify(r.png(html2), seed=101, angle=0.6)
    save_scan_pdf(scan, d / "B_다른포워더견적_스캔.pdf")


def references(r: Renderer, d: Path):
    """참고 서류: 부대비용 요율표(①), 할증료 공지(②), 우리 수수료 메모(⑤⑥)."""
    html = f"""
<div class='page'>
  {carrier_header("CORVANE LINE", "Corvane Line Korea<br>Busan Office", "#163e73")}
  <div style='font-size:20px;font-weight:800;margin:26px 0 4px'>BUSAN EXPORT LOCAL CHARGES TARIFF</div>
  <div class='muted' style='margin-bottom:14px'>Effective from 2026-10-01 until further notice</div>
  <table>
    <tr><th style='width:40%'>Charge</th><th>20'</th><th>40'</th><th>40'HC</th><th>Unit</th></tr>
    <tr><td>THC (Terminal Handling Charge)</td><td class='r'>USD 150</td><td class='r'>USD 200</td><td class='r'>USD 200</td><td class='c'>per CNTR</td></tr>
    <tr><td>DOC Fee (B/L Fee)</td><td class='r' colspan='3'>USD 50</td><td class='c'>per B/L</td></tr>
    <tr><td>Seal Fee</td><td class='r' colspan='3'>USD 10</td><td class='c'>per CNTR</td></tr>
  </table>
  <div class='small muted' style='margin-top:16px;line-height:1.6'>
    Local charges are collected separately from ocean freight unless otherwise stated in the quotation.
  </div>
</div>"""
    r.pdf(html, d / "부대비용요율표_Corvane.pdf")

    html = f"""
<div class='page'>
  {carrier_header("SELVARO SHIPPING", "Selvaro Shipping Korea<br>Customer Advisory", "#1f6f5c")}
  <div style='font-size:20px;font-weight:800;margin:26px 0 4px'>BUNKER ADJUSTMENT FACTOR (BAF) NOTICE</div>
  <div class='muted' style='margin-bottom:14px'>Applicable to cargo loaded from 2026-10-01 to 2026-10-31</div>
  <div style='line-height:1.7;margin-bottom:14px'>Dear Valued Customers,<br>
    Please be advised of the monthly BAF for October 2026 as below.</div>
  <table>
    <tr><th style='width:46%'>Trade</th><th>20'</th><th>40' / 40'HC</th></tr>
    <tr><td>Asia → North Europe</td><td class='r'>USD 210</td><td class='r'>USD 420</td></tr>
    <tr><td>Asia → US West Coast</td><td class='r'>USD 190</td><td class='r'>USD 380</td></tr>
    <tr><td>Intra Asia</td><td class='r'>USD 45</td><td class='r'>USD 90</td></tr>
  </table>
  <div class='small muted' style='margin-top:16px'>Thank you for your continued support.</div>
</div>"""
    r.pdf(html, d / "할증료공지_Selvaro_2026-10.pdf")

    (d / "우리수수료_메모.txt").write_text(
        "견적 회신 시 우리 수수료(Handling Charge): USD 50 / 건\n", encoding="utf-8")


PROBLEMS = [
    ("01_인상이유_부산-로테르담", q01),
    ("02_인상이유_부산-함부르크", q02),
    ("03_인상이유_부산-롱비치", q03),
    ("04_인상이유_부산-싱가포르(대조)", q04),
    ("05_견적회신_부산-로테르담", q05),
    ("06_견적회신_부산-하이퐁", q06),
    ("07_견적회신_부산-LA", q07),
    ("08_재협상_부산-로테르담", q08),
    ("09_재협상_부산-앤트워프", q09),
    ("10_재협상_부산-제벨알리", q10),
]


def main():
    OUT.mkdir(exist_ok=True)
    r = Renderer()
    try:
        for name, fn in PROBLEMS:
            d = OUT / name
            d.mkdir(exist_ok=True)
            fn(r, d)
            print("made", name)
        ref = OUT / "00_참고서류"
        ref.mkdir(exist_ok=True)
        references(r, ref)
        print("made references")
    finally:
        r.close()


if __name__ == "__main__":
    main()
