"""만든 서류를 다시 열어, 4.2 설계의 숫자·문구가 그대로 들어 있는지 확인한다.

글자 PDF·엑셀·메일은 글자를 꺼내 비교한다. 사진·스캔은 그림이라 여기서는
원본 HTML에 넣은 문구 대신 파일이 있는지만 본다(OCR 확인은 2단계에서 한다).
"""
import email
import sys
from email import policy
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

S = Path(__file__).resolve().parents[2] / "samples"


def pdf_text(p):
    with pdfplumber.open(p) as pdf:
        return "\n".join((pg.extract_text() or "") for pg in pdf.pages)


def xlsx_text(p):
    wb = load_workbook(p, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def eml_parts(p):
    m = email.message_from_bytes(p.read_bytes(), policy=policy.default)
    body = m.get_body(preferencelist=("plain",)).get_content()
    atts = [a.get_filename() for a in m.iter_attachments()]
    return body, atts


def latest_body(body):
    for marker in ("-----Original Message-----", "\nFrom: "):
        if marker in body:
            body = body.split(marker)[0]
    return body


CHECKS = [
    # (파일, 꺼내는 법, 들어 있어야 할 문구들, 들어 있으면 안 되는 문구들)
    ("02_인상이유_부산-함부르크/A_지난견적_8월.xlsx", "xlsx",
     ["All-in", "2000", "40'HC", "2026-08", "기본운임(O/F)·유류할증료(BAF) 포함", "터미널 조작료(THC)·서류비 별도"], []),
    ("02_인상이유_부산-함부르크/B_선사견적_10월.pdf", "pdf",
     ["1,650", "as per monthly notice", "as per local tariff", "HAMBURG", "40'HC"], []),
    ("04_인상이유_부산-싱가포르(대조)/A_지난견적_9월.pdf", "pdf", ["USD 520", "THC 포함", "SINGAPORE"], []),
    ("04_인상이유_부산-싱가포르(대조)/B_선사운임_10월.xlsx", "xlsx", ["580", "THC 포함", "SGSIN", "2026-10-01"], []),
    ("05_견적회신_부산-로테르담/B_선사운임표.pdf", "pdf",
     ["1,700", "USD 200", "TBA", "Valid till 2026-10-31", "ROTTERDAM"], []),
    ("06_견적회신_부산-하이퐁/B_선사운임.xlsx", "xlsx", ["HAIPHONG", "480", "200", "CY/CY", "2026-11-15"], []),
    ("07_견적회신_부산-LA/B_첨부_Talvik_USWC_Nov2026.xlsx", "xlsx",
     ["LOS ANGELES", "LONG BEACH", "OAKLAND", "SEATTLE", "VANCOUVER", "N/A"], []),
    ("08_재협상_부산-로테르담/A_우리견적.xlsx", "xlsx", ["1700", "200", "350", "50", "2300", "공장"], []),
    ("10_재협상_부산-제벨알리/A_우리견적.pdf", "pdf",
     ["USD 3,400", "USD 400", "USD 50", "USD 3,850", "1,400원", "40' × 2", "JEBEL ALI"], []),
    ("00_참고서류/부대비용요율표_Corvane.pdf", "pdf", ["USD 150", "USD 200", "USD 50", "THC"], []),
    ("00_참고서류/할증료공지_Selvaro_2026-10.pdf", "pdf", ["USD 420", "USD 380", "USD 90", "North Europe"], []),
    ("00_참고서류/우리수수료_메모.txt", "txt", ["USD 50"], []),
]

EML_CHECKS = [
    # (파일, 최근 본문에 있어야 할 것, 최근 본문에 없어야 할 것, 전체에 있어야 할 것, 첨부)
    ("01_인상이유_부산-로테르담/B_선사회신_10월.eml",
     ["O/F USD 1,700", "CY/CY", "Local charges as per tariff"], ["1,800"], ["-----Original Message-----"], []),
    ("03_인상이유_부산-롱비치/B_선사회신_10월.eml", ["USD 1,760", "THC excl.", "Long Beach"], [], [], []),
    ("05_견적회신_부산-로테르담/A_화주문의.eml", ["40ft 1대", "11월 둘째 주", "도착항 비용까지"], [], [], []),
    ("06_견적회신_부산-하이퐁/A_화주문의_스레드.eml",
     ["하이퐁 40ft 1대", "하노이 공장 문 앞까지", "11월 3일"], ["520"], ["CY 기준 40ft 1대 USD 520"], []),
    ("07_견적회신_부산-LA/B_선사회신.eml", ["11월 운임 첨부"], [], [], ["Talvik_USWC_Nov2026.xlsx"]),
    ("09_재협상_부산-앤트워프/A_우리견적_스레드.eml",
     ["USD 1,980", "11-30"], ["1,850"], ["USD 1,850"], []),
]

IMAGE_FILES = [
    "01_인상이유_부산-로테르담/A_지난견적_9월_인쇄원본.pdf",
    "01_인상이유_부산-로테르담/A_지난견적_9월_사진(합성).jpg",
    "03_인상이유_부산-롱비치/A_지난견적_9월_스캔.pdf",
    "07_견적회신_부산-LA/A_선적요청서_스캔.pdf",
    "08_재협상_부산-로테르담/B_다른포워더견적_인쇄원본.pdf",
    "08_재협상_부산-로테르담/B_다른포워더견적_사진(합성).jpg",
    "09_재협상_부산-앤트워프/B_다른포워더견적_인쇄원본.pdf",
    "09_재협상_부산-앤트워프/B_다른포워더견적_사진(합성).jpg",
    "10_재협상_부산-제벨알리/B_다른포워더견적_스캔.pdf",
]


def main():
    fails = []
    for rel, kind, must, must_not in CHECKS:
        p = S / rel
        text = {"pdf": pdf_text, "xlsx": xlsx_text, "txt": lambda q: q.read_text(encoding="utf-8")}[kind](p)
        for m in must:
            if m not in text:
                fails.append(f"{rel}: '{m}' 없음")
        for m in must_not:
            if m in text:
                fails.append(f"{rel}: '{m}'이 있으면 안 됨")
    for rel, must, must_not, whole, atts in EML_CHECKS:
        body, found_atts = eml_parts(S / rel)
        latest = latest_body(body)
        for m in must:
            if m not in latest:
                fails.append(f"{rel}: 최근 본문에 '{m}' 없음")
        for m in must_not:
            if m in latest:
                fails.append(f"{rel}: 최근 본문에 '{m}'이 있으면 안 됨")
        for m in whole:
            if m not in body:
                fails.append(f"{rel}: 인용 부분에 '{m}' 없음")
        if found_atts != atts:
            fails.append(f"{rel}: 첨부 {found_atts} (기대 {atts})")
    for rel in IMAGE_FILES:
        if not (S / rel).exists():
            fails.append(f"{rel}: 파일 없음")

    # 형식별 개수 (문제 서류 20개)
    probs = [d for d in S.iterdir() if d.is_dir() and d.name[:2].isdigit() and d.name[:2] != "00"]
    count = {"eml": 0, "글자PDF": 0, "엑셀": 0, "사진": 0, "스캔": 0}
    for d in probs:
        for f in d.iterdir():
            n = f.name
            if n.endswith(".eml"):
                count["eml"] += 1
            elif "스캔" in n:
                count["스캔"] += 1
            elif "사진(합성)" in n:
                count["사진"] += 1
            elif n.endswith(".pdf") and "인쇄원본" not in n:
                count["글자PDF"] += 1
            elif n.endswith(".xlsx") and "첨부" not in n:
                count["엑셀"] += 1
    expected = {"eml": 6, "글자PDF": 4, "엑셀": 4, "사진": 3, "스캔": 3}
    if count != expected:
        fails.append(f"형식별 개수 {count} (기대 {expected})")

    if fails:
        print("\n".join(fails))
        sys.exit(1)
    print("모두 통과:", count, "합계", sum(count.values()))


if __name__ == "__main__":
    main()
